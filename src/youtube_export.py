from docx import Document

def save_notes_docx(notes, filename):

    doc = Document()

    doc.add_heading(
        "YouTube Study Notes",
        level=1
    )

    doc.add_paragraph(notes)

    doc.save(filename)